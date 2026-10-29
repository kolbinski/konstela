# -*- coding: utf-8 -*-
"""Projekty pism intencyjnych jako docx: dla MKiDN i dla wydawcy.

Oba mieszczą się na jednej stronie A4 i są pomyślane do swobodnej edycji
po stronie adresata. Uruchomienie: python3 skrypty/docx-listy-intencyjne.py
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

KATALOG = 'attachments'


def nowy_dokument():
    d = Document()
    s = d.sections[0]
    s.top_margin = Cm(1.5)
    s.bottom_margin = Cm(1.5)
    s.left_margin = Cm(2)
    s.right_margin = Cm(2)
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.0
    return d


def akapit(d, tekst='', *, align=None, przed=0, po=6, rozmiar=None):
    """Tekst między gwiazdkami jest pogrubiony."""
    par = d.add_paragraph()
    par.paragraph_format.space_before = Pt(przed)
    par.paragraph_format.space_after = Pt(po)
    if align:
        par.alignment = align
    for i, czesc in enumerate(tekst.split('**')):
        if not czesc:
            continue
        r = par.add_run(czesc)
        if i % 2 == 1:
            r.bold = True
        if rozmiar:
            r.font.size = Pt(rozmiar)
    return par


def nota_projektu(d, tresc):
    akapit(d, tresc, po=10, rozmiar=8.5)


def stopka_podpisu(d, pieczec, nota_o_podpisie=False):
    akapit(d, pieczec, przed=10, po=20)
    akapit(d, '..........................................')
    akapit(d, 'podpis, imię i nazwisko, stanowisko', po=14)
    if nota_o_podpisie:
        akapit(d,
               'Pismo może zostać opatrzone kwalifikowanym podpisem elektronicznym zamiast '
               'podpisu odręcznego i pieczęci.', rozmiar=8.5)


def mkidn():
    d = nowy_dokument()
    nota_projektu(d, '**PROJEKT POMOCNICZY, DO SWOBODNEJ ZMIANY LUB ODRZUCENIA.** Treść poniższa '
                     'jest propozycją wnioskodawcy; o formie i zakresie ewentualnego stanowiska '
                     'rozstrzyga Ministerstwo.')
    akapit(d, '[nagłówek Ministerstwa Kultury i Dziedzictwa Narodowego]', przed=6)
    akapit(d, 'Warszawa, dnia ..........', przed=10)
    akapit(d, 'Znak sprawy: ..........', po=12)
    akapit(d, 'Pan Krzysztof Olbiński')
    akapit(d, 'ul. 11 Listopada 12/1')
    akapit(d, '98-200 Sieradz', po=14)
    akapit(d, '**LIST INTENCYJNY / STANOWISKO DOTYCZĄCE KONCEPCJI NAGRODY KONSTELA**',
           align=WD_ALIGN_PARAGRAPH.CENTER, po=0)
    akapit(d, '(tytuł do wyboru przez Ministerstwo, zależnie od przyjętej formy)',
           align=WD_ALIGN_PARAGRAPH.CENTER, po=14, rozmiar=8.5)
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    akapit(d, 'Ministerstwo Kultury i Dziedzictwa Narodowego otrzymało opis koncepcji Nagrody '
              'Konstela - cyklicznej nagrody przyznawanej osobom prowadzącym pogłębione rozmowy '
              'z systemami sztucznej inteligencji - oraz towarzyszących jej publikacji: antologii '
              'najwyżej ocenionych rozmów każdej edycji, wydawanej nieodpłatnie, i publicznego '
              'Archiwum szerszego zbioru rozmów.', align=J)
    akapit(d, 'Ministerstwo odnotowuje, że tematyka przedstawionego przedsięwzięcia dotyczy obszaru '
              'upowszechniania kultury i czytelnictwa oraz dokumentowania dorobku polskojęzycznego. '
              'Zgodnie z przedstawionymi założeniami powstający zbiór ma być zapisem tego, jak '
              'użytkownicy języka polskiego posługują się nowym narzędziem w okresie jego wchodzenia '
              'do powszechnego użytku, a antologia i Archiwum mają być udostępniane bezpłatnie, '
              'bez rejestracji i bez ograniczeń dostępu.', align=J)
    akapit(d, 'Niniejsze stanowisko odnosi się wyłącznie do tematyki przedsięwzięcia w odniesieniu '
              'do obszaru działania Ministerstwa i opiera się na opisie przedstawionym przez '
              'adresata. Nie zawiera oceny jakości plebiscytu, jego regulaminu ani przyjętych '
              'kryteriów oceny, nie stanowi rekomendacji organizatora i nie jest zapewnieniem, '
              'że przedsięwzięcie zostanie zrealizowane.', align=J)
    akapit(d, 'Niniejsze pismo nie stanowi przyznania patronatu honorowego, zobowiązania finansowego '
              'ani zobowiązania do podjęcia innych działań po stronie Ministerstwa. Nie przesądza ono '
              'w szczególności o sposobie rozpoznania ewentualnego wniosku o dofinansowanie złożonego '
              'w ramach Programów Ministra ani wniosku o patronat honorowy; każdy taki wniosek '
              'zostanie oceniony na zasadach ogólnych, określonych w regulaminie właściwego programu '
              'lub w trybie patronatowym.', align=J)
    akapit(d, 'Adresat może posługiwać się niniejszym pismem w kontaktach z potencjalnymi partnerami '
              'i fundatorami Nagrody oraz publikować jego treść w całości, bez cytowania fragmentów '
              'w oderwaniu od pozostałych postanowień.', align=J)
    stopka_podpisu(d, '(pieczęć urzędowa)', nota_o_podpisie=True)
    d.save(KATALOG + '/mk-list-intencyjny.docx')


def wydawca():
    d = nowy_dokument()
    nota_projektu(d, '**PROJEKT POMOCNICZY, DO SWOBODNEJ ZMIANY LUB ODRZUCENIA.** Treść poniższa '
                     'jest propozycją organizatora Nagrody; o formie i zakresie ewentualnego pisma '
                     'rozstrzyga wydawnictwo.')
    akapit(d, '[nagłówek wydawnictwa]', przed=6)
    akapit(d, '.........................., dnia ..........', przed=10, po=12)
    akapit(d, 'Pan Krzysztof Olbiński')
    akapit(d, 'ul. 11 Listopada 12/1')
    akapit(d, '98-200 Sieradz', po=14)
    akapit(d, '**OŚWIADCZENIE O WSTĘPNYM ZAINTERESOWANIU**',
           align=WD_ALIGN_PARAGRAPH.CENTER, po=14)
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    akapit(d, 'Wydawnictwo zapoznało się z przedstawioną propozycją wydawniczą antologii Nagrody '
              'Konstela - corocznego tomu zawierającego dziesięć najwyżej ocenionych rozmów '
              'człowieka ze sztuczną inteligencją, publikowanych w całości wraz z komentarzem '
              'jury.', align=J)
    akapit(d, 'Wydawnictwo stwierdza, że tematyka propozycji mieści się w obszarze jego '
              'zainteresowań wydawniczych, i deklaruje **wstępne zainteresowanie oraz gotowość '
              'rozważenia wydania tomu** po otrzymaniu kompletnego materiału.', align=J)
    akapit(d, 'Niniejsze oświadczenie nie przesądza o decyzji wydawniczej. Wydawnictwo podejmie ją '
              'po zapoznaniu się z pełnym maszynopisem, zgodnie z własnym trybem kwalifikowania '
              'tytułów, i nie jest zobowiązane do wydania tomu ani do zawarcia jakiejkolwiek umowy. '
              'Oświadczenie nie zawiera zobowiązania finansowego, nie stanowi umowy wydawniczej ani '
              'przyrzeczenia jej zawarcia i nie ustala warunków ewentualnej współpracy.', align=J)
    akapit(d, 'Adresat może posługiwać się niniejszym pismem w kontaktach z potencjalnymi partnerami '
              'i fundatorami Nagrody oraz publikować jego treść w całości, bez cytowania fragmentów '
              'w oderwaniu od pozostałych postanowień.', align=J)
    stopka_podpisu(d, '(pieczęć wydawnictwa)')
    d.save(KATALOG + '/antologia-list-intencyjny.docx')


if __name__ == '__main__':
    mkidn()
    wydawca()
    print('Zapisane:', KATALOG + '/mk-list-intencyjny.docx', 'oraz', KATALOG + '/antologia-list-intencyjny.docx')
